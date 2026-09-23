from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).parent
DOC = ROOT / 'signature-management-use-case-specifications.docx'
IMG = Path(r'C:\Users\Admin\Downloads\Mock Up\FE12-Signature Management')

screens = [
('3.1.1.1 List Signature Screen','Image Jun 22, 2026, 08_59_24 PM (1).png','UC-14 List Electronic Signature',
 'Displays the current user’s electronic signatures. The user can search, filter, paginate, open a signature, or start creation.',[
 ('Search and filters',None),
 ('Search signatures','Text input; searches names in the loaded list. No maximum length is defined in the current frontend.'),
 ('Type','Dropdown; filters by electronic signature type. Initial value: All.'),
 ('Status','Dropdown; filters by signature status. Initial value: All.'),
 ('Filters','Button shown in the mockup for filter controls; the current page uses direct type and status filters.'),
 ('Refresh','Reloads the signature list from the backend.'),
 ('Signature table',None),
 ('Signature Name','Text and image preview for each signature. The API projection uses signatureName and fileUrl.'),
 ('Type','ElectronicSignatureType value for each signature.'),
 ('Used In','Usage category shown in the mockup; not returned by the current list API.'),
 ('Status','ElectronicStatus value for each signature.'),
 ('Updated At','Date shown in the mockup; the current list API returns uploadAt from createdAt, not updatedAt.'),
 ('Actions','Opens actions for a selected signature, including viewing its details.'),
 ('Pagination','Moves between pages of the client-side filtered list.'),
 ('Page size','Selects the number of rows shown per page; the current default is 10.'),
 ('New Signature','Opens the Create Signature screen.')]),
('3.1.1.2 Create Signature Screen','Image Jun 22, 2026, 08_59_24 PM (2).png','UC-16 Create Electronic Signature',
 'Collects signature details and an image, generates an RSA key pair, and submits a new signature for the current user.',[
 ('Signature Information',None),
 ('Signature Name','Required text input; stored as electronicSignatureName. The current request defines no maximum length.'),
 ('Signature Type','Required dropdown; stored as ElectronicSignatureType. The mockup offers Drawn, Uploaded, and Typed modes.'),
 ('Used In','Required dropdown in the mockup; not defined in the current create request or signature entity.'),
 ('Description','Optional text area in the mockup; not defined in the current create request or signature entity.'),
 ('Status','Required dropdown; stored as ElectronicStatus. Initial mockup value: Active.'),
 ('Set as Default','Boolean toggle; stored as isDefault. Initial mockup value: enabled.'),
 ('Access Scope','Dropdown in the mockup; not defined in the current create request or signature entity.'),
 ('Last Modified','Read-only date displayed in the mockup; not entered by the user during creation.'),
 ('Signature Canvas and key',None),
 ('Draw / Upload / Type','Modes for producing a signature image; the submitted request must contain multipartFile.'),
 ('PIN and Confirm PIN','Browser-side values used to protect the generated private key; not saved in the signature request.'),
 ('Generate Key','Requests RSA public/private key information and a certificate.'),
 ('Clear / Undo','Canvas controls shown in the mockup for editing the drawn image.'),
 ('Preview','Shows the image before submission.'),
 ('Save Signature','Submits name, type, status, default flag, public key, key code, certificate, and image.'),
 ('Cancel','Leaves the create screen without submitting.')]),
('3.1.1.3 View Signature Screen','Image Jun 22, 2026, 08_59_24 PM (3).png','UC-15 View Electronic Signature',
 'Displays the selected user-owned electronic signature and offers navigation to editing.',[
 ('Signature Information',None),
 ('Signature Name','Read-only String returned as electronicSignatureName.'),
 ('Signature Type','Read-only ElectronicSignatureType value.'),
 ('Used In','Read-only usage category shown in the mockup; not returned by the current detail API.'),
 ('Description','Read-only text shown in the mockup; not returned by the current detail API.'),
 ('Status','Read-only ElectronicStatus returned as electronicSignatureStatus.'),
 ('Default Signature','Read-only Boolean returned as isDefault.'),
 ('Access Scope','Read-only access setting shown in the mockup; not returned by the current detail API.'),
 ('Last Updated','Read-only date shown in the mockup; the current detail API returns createAt instead.'),
 ('Signature Image','Displays the image from signatureUrl.'),
 ('Usage Summary',None),
 ('Contract Templates / Approval Workflows / Saved Drafts / Recent Activity','Mockup usage metrics; not returned by the current signature detail API.'),
 ('Edit Signature','Opens the Update Signature screen for the selected ID.')]),
('3.1.1.4 Update Signature Screen','Image Jun 22, 2026, 08_59_24 PM (4).png','UC-17 Update Electronic Signature',
 'Loads an existing signature, lets its owner edit supported fields or replace the image, and saves the changes.',[
 ('Signature Information',None),
 ('Signature Name','Required text input; updates electronicSignatureName. Blank names are rejected by the frontend.'),
 ('Signature Type','Dropdown; updates electronicSignatureType.'),
 ('Used In','Dropdown shown in the mockup; not defined in UpdateElectronicSignatureRequest.'),
 ('Description','Text area shown in the mockup; not defined in UpdateElectronicSignatureRequest.'),
 ('Status','Dropdown; updates electronicStatus.'),
 ('Set as Default','Boolean toggle; updates isDefault.'),
 ('Access Scope','Dropdown shown in the mockup; not defined in UpdateElectronicSignatureRequest.'),
 ('Last Modified','Read-only date shown in the mockup; updatedAt is set by the backend when saving.'),
 ('Edit Signature Canvas',None),
 ('Draw / Upload / Type','Modes shown for replacing the image; a new multipartFile is optional in the update API.'),
 ('Clear / Reset / Preview','Controls shown in the mockup for editing and checking the image before saving.'),
 ('Save Changes','Submits the edited fields and optional image; the backend checks signature ownership.'),
 ('Cancel','Leaves the update screen without saving.')])
]

doc=Document(DOC)
doc.add_page_break()
p=doc.add_paragraph('Signature Management Screen and Field Specifications','Heading 1')
for run in p.runs: run.font.color.rgb=RGBColor(0,0,0)
doc.add_paragraph('These screens correspond to the four electronic signature use cases above. Field descriptions identify mockup-only fields that are not present in the current API.')

def border(cell):
 tcPr=cell._tc.get_or_add_tcPr(); bd=OxmlElement('w:tcBorders')
 for side in ('top','left','bottom','right'):
  edge=OxmlElement('w:'+side); edge.set(qn('w:val'),'single'); edge.set(qn('w:sz'),'5'); edge.set(qn('w:color'),'000000'); bd.append(edge)
 tcPr.append(bd)
def set_cell(cell,text,bold=False):
 cell.text=''; p=cell.paragraphs[0]; p.paragraph_format.space_after=Pt(0)
 r=p.add_run(text); r.bold=bold; r.font.name='Times New Roman'; r.font.size=Pt(9); r.font.color.rgb=RGBColor(0,0,0)
 border(cell)
def group(table,name):
 row=table.add_row(); merged=row.cells[0].merge(row.cells[1]); set_cell(merged,name,True)

for index,(title,image_name,uc,description,fields) in enumerate(screens):
 if index: doc.add_page_break()
 h=doc.add_paragraph(title,'Heading 2')
 for run in h.runs: run.font.name='Times New Roman'; run.font.color.rgb=RGBColor(0,0,0)
 p=doc.add_paragraph(); p.add_run('UI layout').bold=True
 image_path=IMG/image_name
 assert image_path.exists(),image_path
 pic=doc.add_paragraph(); pic.alignment=WD_ALIGN_PARAGRAPH.CENTER
 pic.add_run().add_picture(str(image_path),width=Inches(6.8))
 p=doc.add_paragraph(); p.add_run('Function and related use case').bold=True
 doc.add_paragraph(f'{description} Related use case: {uc}.')
 p=doc.add_paragraph(); p.add_run('Screen components and fields').bold=True
 table=doc.add_table(rows=1,cols=2); table.autofit=False
 table.columns[0].width=Inches(2.2); table.columns[1].width=Inches(5.0)
 set_cell(table.rows[0].cells[0],'Field Name',True)
 set_cell(table.rows[0].cells[1],'Description',True)
 table.rows[0]._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
 for name,desc in fields:
  if desc is None: group(table,name)
  else:
   cells=table.add_row().cells; set_cell(cells[0],name); set_cell(cells[1],desc)

doc.save(DOC)
print(DOC)
