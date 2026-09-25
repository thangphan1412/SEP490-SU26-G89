from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

OUT=Path(__file__).with_name('authentication-use-case-and-screen-specifications.docx')
IMG=Path(r'C:\Users\Admin\Downloads\Mock Up\FE01-Authentication')

uses=[
('UC-01 Login','Registered User, External Partner, Employee, Administrative Staff, Accountant Staff, HOD, CEO','None',
 'As a user, I want to sign in so that I can access functions allowed by my role.',
 ['The user has an account.','The sign-in screen is available.'],
 ['A JWT is returned and stored in the browser.','The user is routed to the home page or to signature-key creation.'],
 ['The user opens the Sign In screen.','The user enters an email address and password.','The user selects Sign In.','The system authenticates the credentials and returns a JWT, role, name, department, and signature-key status.','The browser stores the returned values.','The browser navigates to signature-key creation if no active key exists; otherwise it navigates to the home page.'],
 [('A1 Invalid credentials',['The authentication request fails.','The user remains on the sign-in screen.']),('A2 No active signature key',['Authentication succeeds but hasSignatureKey is false.','The browser opens Create Signature instead of the home page.'])]),
('UC-04 Change Password','Registered User, External Partner, Employee, Administrative Staff, Accountant Staff, HOD, CEO','None',
 'As an authenticated user, I want to change my password to protect my account.',
 ['The user is authenticated.','The user knows the current password.'],
 ['The account password is updated when validation succeeds.'],
 ['The user opens Change Password.','The user enters the current password, a new password, and confirmation.','The user selects Update Password.','The system verifies the current password and validates the new password.','The system updates the password and displays success.'],
 [('A1 Incorrect current password',['The system rejects the change.','The screen displays an error.']),('A2 Invalid or mismatched new password',['The system rejects the change.','The user corrects the new password or confirmation.'])]),
('UC-05 Forgot Password','Registered User, External Partner, Employee, Administrative Staff, Accountant Staff, HOD, CEO','Gmail',
 'As a user, I want to request password recovery using my email address.',
 ['The user can access the Forgot Password screen.','An email address is available for the account.'],
 ['The system stores a five-minute OTP and sends it through email.'],
 ['The user opens Forgot Password from Sign In.','The user enters an email address.','The user selects Send Reset Link.','The system generates and emails an OTP.','The screen continues to the reset step when the request succeeds.'],
 [('A1 Request fails',['The backend cannot process the email request.','The screen remains on the recovery step.'])]),
('UC-06 Reset Password','Registered User, External Partner, Employee, Administrative Staff, Accountant Staff, HOD, CEO','Gmail',
 'As a user, I want to reset my password after recovery verification so that I can sign in again.',
 ['The user has initiated password recovery.','The user has the emailed OTP required by the current API.'],
 ['The password is replaced when the OTP and new password are valid.'],
 ['The user opens the Reset Password screen.','The user provides an email address and OTP.','The user enters and confirms a new password.','The user selects Reset Password.','The system validates the OTP and password and updates the account.','The user returns to Sign In.'],
 [('A1 Invalid or expired OTP',['The system rejects the reset.','The user requests or enters a valid OTP.']),('A2 Invalid or mismatched password',['The system rejects the reset.','The user corrects the password fields.'])])
]

screens=[
('3.1.1.1 Sign In Screen','Image Jun 11, 2026, 02_39_21 PM (7).png','UC-01 Login','Collects account credentials and starts authentication.',[
 ('Credentials',None),('Email address','Email input; required to identify the account. No length limit is defined in the current login form.'),('Password','Masked text input; required for authentication. The eye button changes visibility.'),('Remember me','Checkbox shown in the mockup and frontend; current login code does not use this value when storing the token.'),('Sign In','Submits email and password to POST /auth/login.'),('Navigation',None),('Forgot password?','Opens the Forgot Password screen.'),('Sign up','Opens the registration mockup; the current router does not define /register.')],
 'The current frontend sends credentials and stores a JWT, role, name, and department. It routes users without an active signature key to Create Signature.'),
('3.1.1.2 Access Granted Screen','Image Jun 11, 2026, 02_39_21 PM (6).png','UC-01 Login','Shows successful sign-in and the user’s role before entering the dashboard.',[
 ('Account information',None),('User name','Read-only display of the signed-in user name.'),('Role','Read-only display of the assigned role.'),('Permissions','Read-only list of permitted actions shown in the mockup.'),('Go to Dashboard','Opens the dashboard after successful authentication.')],
 'This is a mockup success screen, not a separate use case. The current login frontend navigates directly to /home_page or Create Signature; it does not render this intermediate screen.'),
('3.1.1.3 Change Password Screen','Image Jun 11, 2026, 02_39_21 PM (4).png','UC-04 Change Password','Allows an authenticated user to replace the account password.',[
 ('Password fields',None),('Current password','Masked input sent as oldPassword.'),('New password','Masked input sent as newPassword. Backend requires at least 9 characters, including lowercase, uppercase, digit, and special character.'),('Confirm new password','Masked input for matching newPasswordConfirm.'),('Password strength','Visual indicator shown in the mockup; backend validation uses the rules above.'),('Update Password','Submits the password change request.'),('Show or hide password','Toggles visibility of password inputs.')],
 'The current implementation places password changing in User Profile Management rather than on this standalone mockup page.'),
('3.1.1.4 Forgot Password Screen','Image Jun 11, 2026, 02_39_21 PM (8).png','UC-05 Forgot Password','Starts account recovery using an email address.',[
 ('Recovery form',None),('Email address','Email input sent to POST /auth/forgot-password.'),('Send Reset Link','Starts the recovery request. The mockup says link, while the current reset API requires an OTP.'),('Back to sign in','Returns to /login.')],
 'The current frontend navigates directly to /reset-password after a successful request. Its rendered heading and button text differ from the mockup.'),
('3.1.1.5 Reset Password Screen','Image Jun 11, 2026, 02_39_21 PM (9).png','UC-06 Reset Password','Collects the information needed to replace a forgotten password.',[
 ('Reset form',None),('Email address','Required by the current reset API, although omitted from the mockup.'),('OTP','Required by the current reset API, although omitted from the mockup.'),('New password','Masked input; backend requires at least 9 characters with lowercase, uppercase, digit, and special character.'),('Confirm new password','Masked input; must match the new password.'),('Password strength','Visual indicator shown in the mockup.'),('Reset Password','Submits email, OTP, newPassword, and newPasswordConfirm.'),('Back to sign in','Returns to /login.')],
 'The current ResetPasswordForm contains email and OTP inputs but labels both as email; its heading and button text also differ from the mockup.'),
('3.1.1.6 Create Account Screen','Image Jun 11, 2026, 02_39_21 PM (5).png','Mockup only — no registration use case in the supplied diagrams','Collects information for a new account in the mockup.',[
 ('Account information',None),('Full name','Text input shown in the mockup. The backend request instead has separate firstName and lastName fields.'),('Email address','Email input; maps to RegisterRequest.email.'),('Company name','Text input shown in the mockup; absent from RegisterRequest.'),('Role','Dropdown shown in the mockup; absent from RegisterRequest.'),('Password','Masked input; maps to RegisterRequest.password.'),('Confirm password','Masked input shown in the mockup; absent from RegisterRequest.'),('Terms agreement','Checkbox shown in the mockup; absent from RegisterRequest.'),('Create Account','Submits registration information in the mockup.'),('Sign in','Returns to the Sign In screen.')],
 'The backend has POST /register, but the current RegisterPage is commented out and /register is not present in AppRouter. This screen is a proposed design, not an active frontend flow.')
]

doc=Document(); sec=doc.sections[0]
sec.top_margin=sec.bottom_margin=Inches(.65); sec.left_margin=sec.right_margin=Inches(.65)
style=doc.styles['Normal']; style.font.name='Times New Roman'; style.font.size=Pt(10); style.font.color.rgb=RGBColor(0,0,0)
title=doc.add_paragraph('Authentication Use Case and Screen Specifications','Title')
for r in title.runs: r.font.name='Times New Roman'; r.font.size=Pt(16); r.font.color.rgb=RGBColor(0,0,0)
doc.add_paragraph('This document covers four authentication use cases and six supplied screen mockups. Differences between the mockups and current code are stated in each screen section.')

def border(cell):
 tcPr=cell._tc.get_or_add_tcPr(); bd=OxmlElement('w:tcBorders')
 for side in ('top','left','bottom','right'):
  e=OxmlElement('w:'+side); e.set(qn('w:val'),'single'); e.set(qn('w:sz'),'5'); e.set(qn('w:color'),'000000'); bd.append(e)
 tcPr.append(bd)
def put(cell,value,bold=False):
 cell.text=''; p=cell.paragraphs[0]; p.paragraph_format.space_after=Pt(0)
 r=p.add_run(value); r.bold=bold; r.font.name='Times New Roman'; r.font.size=Pt(9.5); r.font.color.rgb=RGBColor(0,0,0); border(cell)
def listcell(cell,items,number=False):
 put(cell,'')
 for n,item in enumerate(items,1):
  p=cell.paragraphs[0] if n==1 else cell.add_paragraph()
  p.paragraph_format.space_after=Pt(2)
  r=p.add_run((f'{n}. ' if number else '• ') + item); r.font.name='Times New Roman'; r.font.size=Pt(9.5); r.font.color.rgb=RGBColor(0,0,0)
def two(table,label,value,mode=None):
 c=table.add_row().cells; put(c[0],label,True); x=c[1].merge(c[3])
 if mode: listcell(x,value,mode=='number')
 else: put(x,value)

for i,(name,primary,secondary,description,pre,post,normal,alts) in enumerate(uses):
 if i: doc.add_page_break()
 p=doc.add_paragraph(name,'Heading 1')
 for r in p.runs: r.font.color.rgb=RGBColor(0,0,0)
 t=doc.add_table(rows=0,cols=4); t.autofit=False
 for col,w in zip(t.columns,[Inches(1.35),Inches(2.0),Inches(1.35),Inches(2.5)]): col.width=w
 c=t.add_row().cells
 for x,v,b in zip(c,['Primary Actors',primary,'Secondary Actors',secondary],[True,False,True,False]): put(x,v,b)
 two(t,'Description',description)
 two(t,'Preconditions',pre,'bullet')
 two(t,'Postconditions',post,'bullet')
 two(t,'Normal Sequence/Flow',normal,'number')
 alt=[]
 for heading,steps in alts:
  alt.append(heading)
  alt.extend(f'    {j}. {step}' for j,step in enumerate(steps,1))
 two(t,'Alternative Sequences/Flows',alt,'bullet')

doc.add_page_break(); h=doc.add_paragraph('Authentication Screen and Field Specifications','Heading 1')
for r in h.runs: r.font.color.rgb=RGBColor(0,0,0)
for i,(heading,image,uc,description,fields,note) in enumerate(screens):
 if i: doc.add_page_break()
 h=doc.add_paragraph(heading,'Heading 2')
 for r in h.runs: r.font.color.rgb=RGBColor(0,0,0)
 p=doc.add_paragraph(); p.add_run('UI layout').bold=True
 path=IMG/image; assert path.exists(),path
 p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.add_run().add_picture(str(path),width=Inches(6.8))
 p=doc.add_paragraph(); p.add_run('Function and related use case').bold=True
 doc.add_paragraph(f'{description} Related use case: {uc}.')
 p=doc.add_paragraph(); p.add_run('Screen components and fields').bold=True
 t=doc.add_table(rows=1,cols=2); t.autofit=False
 t.columns[0].width=Inches(2.15); t.columns[1].width=Inches(5.05)
 put(t.rows[0].cells[0],'Field Name',True); put(t.rows[0].cells[1],'Description',True)
 t.rows[0]._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
 for name,desc in fields:
  c=t.add_row().cells
  if desc is None: put(c[0].merge(c[1]),name,True)
  else: put(c[0],name); put(c[1],desc)
 p=doc.add_paragraph(); p.add_run('Source alignment note: ').bold=True; p.add_run(note)

doc.save(OUT); print(OUT)
