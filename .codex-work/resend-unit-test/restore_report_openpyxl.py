from __future__ import annotations

import json
import os
import subprocess
import sys
import zipfile
from copy import copy
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.cell.cell import MergedCell
from openpyxl.styles import Alignment, Font


ROOT = Path(__file__).resolve().parents[2]
SOURCE = Path(r"C:\Users\acer\Downloads\Report5.1_Unit Test.xlsx")
OUTPUT = ROOT / "outputs" / "01a0a602-3ae5-73c1-8485-d4f67d55a97c" / "Report5.1_Unit_Test.xlsx"
NODE = Path(r"C:\Users\acer\.cache\codex-runtimes\codex-primary-runtime\dependencies\node\bin\node.exe")
EXTRACTOR = Path(__file__).with_name("extract_cases.mjs")
EXECUTED_DATE = datetime(2026, 9, 16)


def load_definitions() -> list[dict]:
    completed = subprocess.run(
        [str(NODE), str(EXTRACTOR)],
        cwd=str(EXTRACTOR.parent),
        check=True,
        capture_output=True,
        text=True,
        encoding="utf-8",
    )
    return json.loads(completed.stdout)


def clear_range(ws, min_row: int, max_row: int, min_col: int, max_col: int) -> None:
    for row in ws.iter_rows(
        min_row=min_row, max_row=max_row, min_col=min_col, max_col=max_col
    ):
        for cell in row:
            if not isinstance(cell, MergedCell):
                cell.value = None


def put(ws, coordinate: str, value) -> None:
    ws[coordinate] = value


def configure_calculation(wb) -> None:
    calculation = getattr(wb, "calculation", None)
    if calculation is None:
        return
    calculation.calcMode = "auto"
    calculation.fullCalcOnLoad = True
    calculation.forceFullCalc = True
    calculation.calcId = 0


def fill_function_sheet(ws, config: dict) -> dict[str, int]:
    compact = config["sheet"] == "getContracts"
    confirm_row = 30 if compact else 31
    result_row = 38 if compact else 43
    passed_row = result_row + 1
    date_row = result_row + 2
    defect_row = result_row + 3
    cases = config["cases"]

    put(ws, "C1", config["module"])
    put(ws, "L1", config["method"])
    put(ws, "C2", "Project team")
    put(ws, "L2", "Codex / Maven")
    put(ws, "C3", config["requirement"])

    clear_range(ws, 8, 42, 2, 5)
    clear_range(ws, 7, 46, 6, 20)

    put(ws, "A8", "Condition")
    put(ws, "B8", "Precondition")
    put(ws, "C8", "Scenario / Input")
    put(ws, "B9", "Setup")
    put(ws, "B10", "Test case")
    put(ws, f"A{confirm_row}", "Confirm")
    put(ws, "B31", "Return")
    put(ws, "C31", "Expected result")
    put(ws, "B32", "Expected")
    put(ws, f"A{result_row}", "Result")
    put(ws, f"B{result_row}", "Type(N : Normal, A : Abnormal, B : Boundary)")
    put(ws, f"B{passed_row}", "Passed/Failed")
    put(ws, f"B{date_row}", "Executed Date")
    put(ws, f"B{defect_row}", "Defect ID")

    type_counts = {kind: sum(1 for case in cases if case[2] == kind) for kind in ("N", "A", "B")}
    put(ws, "A5", len(cases))
    put(ws, "C5", 0)
    put(ws, "F5", 0)
    put(ws, "L5", type_counts["N"])
    put(ws, "M5", type_counts["A"])
    put(ws, "N5", type_counts["B"])
    put(ws, "O5", len(cases))

    put(ws, "C9", "JUnit 5 + Mockito; dependencies mocked")
    for index, (scenario, outcome, case_type) in enumerate(cases):
        column = 6 + index
        put(ws, f"{ws.cell(7, column).column_letter}7", f"UTCID{index + 1:02d}")
        put(ws, f"{ws.cell(9, column).column_letter}9", "O")
        put(ws, f"C{10 + index}", scenario)
        put(ws, f"{ws.cell(10 + index, column).column_letter}{10 + index}", "O")
        put(ws, f"C{32 + index}", outcome)
        put(ws, f"{ws.cell(32 + index, column).column_letter}{32 + index}", "O")
        put(ws, f"{ws.cell(result_row, column).column_letter}{result_row}", case_type)
        put(ws, f"{ws.cell(passed_row, column).column_letter}{passed_row}", "P")
        date_cell = ws.cell(date_row, column)
        date_cell.value = EXECUTED_DATE
        date_cell.number_format = "d/m/yy"
        date_cell.font = Font(name="Arial", size=6, color="000000")

    for cell in ws[7][5:20]:
        cell.font = Font(name="Arial", size=6, bold=True, color="FFFFFF")

    ws["C3"].font = Font(name="Arial", size=7, italic=True, color="000000")
    ws["C3"].alignment = copy(ws["C3"].alignment)
    ws["C3"].alignment = Alignment(
        horizontal=ws["C3"].alignment.horizontal,
        vertical=ws["C3"].alignment.vertical,
        text_rotation=ws["C3"].alignment.text_rotation,
        wrap_text=True,
        shrink_to_fit=ws["C3"].alignment.shrink_to_fit,
        indent=ws["C3"].alignment.indent,
    )
    ws.row_dimensions[3].height = 30

    for row in list(range(9, 25)) + list(range(32, 43)):
        cell = ws.cell(row, 3)
        cell.font = Font(name="Arial", size=7, color="000000")
        cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
        ws.row_dimensions[row].height = 36

    return {
        "passed": len(cases),
        "failed": 0,
        "untested": 0,
        "normal": type_counts["N"],
        "abnormal": type_counts["A"],
        "boundary": type_counts["B"],
        "total": len(cases),
        "result_row": result_row,
        "passed_row": passed_row,
        "date_row": date_row,
    }


def main() -> None:
    if not SOURCE.is_file():
        raise FileNotFoundError(f"Template not found: {SOURCE}")

    definitions = load_definitions()
    wb = load_workbook(SOURCE, data_only=False, keep_links=True)
    original_names = list(wb.sheetnames)
    results: dict[str, dict[str, int]] = {}

    for config in definitions:
        if config["sheet"] not in wb.sheetnames:
            raise KeyError(f"Template sheet missing: {config['sheet']}")
        results[config["sheet"]] = fill_function_sheet(wb[config["sheet"]], config)

    cover = wb["Cover"]
    put(cover, "F5", EXECUTED_DATE)
    put(cover, "F6", 4)
    cover_values = [
        EXECUTED_DATE,
        4,
        "Unit test scope",
        "A",
        "Add executed unit tests for 12 requested screens",
        "Report5.1_Unit Test",
    ]
    for col, value in enumerate(cover_values, start=1):
        cover.cell(14, col).value = value

    method_list = wb["MethodList"]
    put(
        method_list,
        "C6",
        "Java 21; Maven 3.9.2; JUnit 5; Mockito; Spring Boot 4.0.6; Windows 11; source snapshot 1164f260.",
    )
    clear_range(method_list, 9, 20, 1, 6)
    for row, config in enumerate(definitions, start=9):
        values = [
            row - 8,
            config["module"],
            f"{config['method']}()",
            config["sheet"],
            config["screen"],
            "JUnit 5; Mockito mocks; authenticated context where required",
        ]
        for col, value in enumerate(values, start=1):
            method_list.cell(row, col).value = value
        for col in (3, 4):
            method_list.cell(row, col).font = Font(name="Arial", size=6, color="000000")

    stats = wb["Statistics"]
    put(stats, "F4", "Codex / Maven")
    put(stats, "F6", EXECUTED_DATE)
    put(
        stats,
        "B7",
        "80/80 tests passed on source snapshot 1164f260. Current working tree compile is blocked by unresolved merge conflicts.",
    )
    clear_range(stats, 12, 22, 1, 9)

    groups = [
        (1, "User Mgmt (4)", ["getAllUsers", "createUser", "getUserById", "updateUser"]),
        (2, "User Profile (2)", ["getMyProfile", "updateMyProfile"]),
        (3, "Company Profile (2)", ["getCompanyProfile", "updateCompanyProfile"]),
        (
            4,
            "Dashboard + Contract (4)",
            ["getOverviewStatistics", "getContracts", "getPendingSignatureDashboard", "getStatisticalReports"],
        ),
    ]
    fields = ["passed", "failed", "untested", "normal", "abnormal", "boundary", "total"]
    totals = {field: 0 for field in fields}
    for row, (number, label, sheets) in enumerate(groups, start=12):
        stats.cell(row, 1).value = number
        stats.cell(row, 2).value = label
        values = []
        for field in fields:
            value = sum(results[name][field] for name in sheets)
            totals[field] += value
            values.append(value)
        for col, value in enumerate(values, start=3):
            stats.cell(row, col).value = value

    stats["B16"] = "Sub total"
    for col, field in enumerate(fields, start=3):
        stats.cell(16, col).value = totals[field]

    percentages = [
        (18, "Test coverage", 100.0),
        (19, "Test successful coverage", 100.0),
        (20, "Normal case", totals["normal"] / totals["total"] * 100),
        (21, "Abnormal case", totals["abnormal"] / totals["total"] * 100),
        (22, "Boundary case", totals["boundary"] / totals["total"] * 100),
    ]
    for row, label, value in percentages:
        stats.cell(row, 2).value = label
        stats.cell(row, 4).value = value
        stats.cell(row, 5).value = "%"

    configure_calculation(wb)
    if wb.sheetnames != original_names:
        raise AssertionError("Template sheet order changed before save")

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUTPUT)

    with zipfile.ZipFile(OUTPUT, "r") as archive:
        bad_member = archive.testzip()
        if bad_member is not None:
            raise AssertionError(f"Corrupt workbook ZIP member: {bad_member}")

    check = load_workbook(OUTPUT, data_only=False, keep_links=True, read_only=False)
    if check.sheetnames != original_names:
        raise AssertionError("Template sheet order changed after save")
    if len(check.sheetnames) != 44:
        raise AssertionError(f"Expected 44 sheets, found {len(check.sheetnames)}")

    expected_totals = {
        "passed": 80,
        "failed": 0,
        "untested": 0,
        "normal": 35,
        "abnormal": 37,
        "boundary": 8,
        "total": 80,
    }
    actual_totals = {
        field: check["Statistics"].cell(16, col).value
        for col, field in enumerate(fields, start=3)
    }
    if actual_totals != expected_totals:
        raise AssertionError(f"Statistics mismatch: {actual_totals}")

    for config in definitions:
        ws = check[config["sheet"]]
        meta = results[config["sheet"]]
        for index, case in enumerate(config["cases"]):
            col = 6 + index
            if ws.cell(7, col).value != f"UTCID{index + 1:02d}":
                raise AssertionError(f"Missing test ID in {config['sheet']} column {col}")
            if ws.cell(meta["result_row"], col).value != case[2]:
                raise AssertionError(f"Wrong test type in {config['sheet']} column {col}")
            if ws.cell(meta["passed_row"], col).value != "P":
                raise AssertionError(f"Wrong result in {config['sheet']} column {col}")

    error_tokens = ("#REF!", "#DIV/0!", "#VALUE!", "#NAME?", "#N/A", "#NUM!", "#NULL!", "#SPILL!", "#CALC!")
    formula_errors = []
    for ws in check.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                if isinstance(cell.value, str) and any(token in cell.value.upper() for token in error_tokens):
                    formula_errors.append(f"{ws.title}!{cell.coordinate}={cell.value}")
    if formula_errors:
        raise AssertionError("Formula errors found: " + "; ".join(formula_errors[:20]))

    check.close()
    print(
        json.dumps(
            {
                "output": str(OUTPUT),
                "size": OUTPUT.stat().st_size,
                "sheet_count": len(original_names),
                "totals": actual_totals,
                "requested_sheets": [item["sheet"] for item in definitions],
            },
            ensure_ascii=False,
            indent=2,
        )
    )


if __name__ == "__main__":
    try:
        main()
    except Exception as exc:
        print(f"ERROR: {exc}", file=sys.stderr)
        raise
